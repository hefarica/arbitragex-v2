#!/usr/bin/env python3
"""
gen-env-mcp.py  —  ArbitrageX v2 MCP secrets bridge (audit / shadow / read-only)

Reads the operator's secrets SSOT (ArbitrageX_Secrets_Config.xlsx) and DERIVES a
gitignored `.env.mcp` so the MCP servers declared in `.mcp.json` get their
`${VAR}` values WITHOUT any hardcoded literal anywhere in the repo.

Doctrine:
  - Secrets live ONLY in the Excel SSOT + the generated `.env.mcp` (gitignored).
  - `.mcp.json` keeps `${VAR}` placeholders — never a real value.
  - Postgres MCP is wired to the READ-ONLY role (ARBX_RO_PASSWORD / arbx_ro).
  - NO foundry PRIVATE_KEY is ever emitted (stays "" in .mcp.json).
  - The console NEVER prints secret values — masked summaries only.

Re-run this whenever the Excel changes:  python scripts/mcp/gen-env-mcp.py

Tunables (env vars, all optional):
  MCP_SECRETS_XLSX  path to the xlsx (default: ~/Downloads/ArbitrageX_Secrets_Config.xlsx)
  PGHOST_RO         Postgres host for the read-only URL (default: localhost  <- SSH tunnel)
  PGPORT_RO         Postgres port (default: 5432)
  PGUSER_RO         Postgres read-only role name (default: arbx_ro)
  REDIS_HOST_RO     Redis host (default: localhost  <- SSH tunnel)
  REDIS_PORT_RO     Redis port (default: 6379)
  EVM_NETWORK       e.g. eth-mainnet, base-mainnet (enables Alchemy EVM_RPC_URL build)
"""
import os, re, sys
from pathlib import Path
from urllib.parse import urlsplit

try:  # Windows consoles default to cp1252; force UTF-8 so output never crashes.
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import openpyxl
except ImportError:
    sys.exit("ERROR: openpyxl not installed.  pip install openpyxl")

REPO_ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(os.environ.get(
    "MCP_SECRETS_XLSX",
    Path.home() / "Downloads" / "ArbitrageX_Secrets_Config.xlsx"))
OUT = REPO_ROOT / ".env.mcp"

VAR_RE = re.compile(r"^[A-Z][A-Z0-9_]{2,}$")


def mask(v: str) -> str:
    if not v:
        return "<empty>"
    s = str(v)
    return f"<len={len(s)} {s[:2]}..>"


def load_secrets(path: Path) -> dict:
    """Merge every (NAME, value) pair found across all sheets where col-A looks
    like an ENV var name. First non-empty wins; later duplicates are ignored."""
    if not path.exists():
        sys.exit(f"ERROR: secrets file not found: {path}\n"
                 f"Set MCP_SECRETS_XLSX or place the xlsx at the default path.")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            name, val = row[0], row[1]
            if name is None or val is None:
                continue
            name, val = str(name).strip(), str(val).strip()
            if VAR_RE.match(name) and val and name not in out:
                out[name] = val
    return out


def main():
    s = load_secrets(XLSX)

    pg_user = os.environ.get("PGUSER_RO", "arbx_ro")
    pg_host = os.environ.get("PGHOST_RO", "localhost")
    pg_port = os.environ.get("PGPORT_RO", "5432")
    rd_host = os.environ.get("REDIS_HOST_RO", "localhost")
    rd_port = os.environ.get("REDIS_PORT_RO", "6379")
    # network for the Alchemy-built EVM RPC: env wins, else an EVM_NETWORK cell
    evm_net = os.environ.get("EVM_NETWORK", "").strip() or (s.get("EVM_NETWORK") or "").strip()

    # DB name from the operator's DATABASE_URL if present, else default.
    pg_db = "arbitragex"
    if s.get("DATABASE_URL"):
        path = urlsplit(s["DATABASE_URL"]).path.lstrip("/")
        if path:
            pg_db = path

    resolved, todos, warnings = {}, [], []

    # --- Postgres (READ-ONLY role) -----------------------------------------
    if s.get("DATABASE_URL_READONLY"):
        resolved["DATABASE_URL_READONLY"] = s["DATABASE_URL_READONLY"]
        src_pg = "cell DATABASE_URL_READONLY (verbatim)"
    elif s.get("ARBX_RO_PASSWORD"):
        pw = s["ARBX_RO_PASSWORD"]
        resolved["DATABASE_URL_READONLY"] = (
            f"postgres://{pg_user}:{pw}@{pg_host}:{pg_port}/{pg_db}")
        src_pg = f"built from ARBX_RO_PASSWORD as {pg_user}@{pg_host}:{pg_port}/{pg_db}"
    else:
        todos.append("DATABASE_URL_READONLY (no ARBX_RO_PASSWORD cell found)")
        src_pg = None

    # --- Redis -------------------------------------------------------------
    if s.get("REDIS_URL_READONLY"):
        resolved["REDIS_URL_READONLY"] = s["REDIS_URL_READONLY"]
        src_rd = "cell REDIS_URL_READONLY (verbatim)"
    elif s.get("REDIS_URL"):
        parts = urlsplit(s["REDIS_URL"])
        userinfo = f"{parts.username}:{parts.password}@" if parts.username else ""
        if parts.password:
            userinfo = f"{parts.username}:{parts.password}@"
        tail = parts.path or ""
        resolved["REDIS_URL_READONLY"] = f"redis://{userinfo}{rd_host}:{rd_port}{tail}"
        src_rd = f"built from REDIS_URL, host->{rd_host}:{rd_port}"
        if not parts.password:
            warnings.append(
                "REDIS has NO auth/ACL. server-redis exposes WRITE tools (set/delete). "
                "Per CLAUDE.md §33.2 create a read-only ACL user on the VPS and add a "
                "'REDIS_URL_READONLY' row to the Excel:\n"
                "      ACL SETUSER arbx_ro on >STRONGPASS ~* +@read -@write -@dangerous\n"
                "      REDIS_URL_READONLY=redis://arbx_ro:STRONGPASS@HOST:PORT")
    else:
        todos.append("REDIS_URL_READONLY (no REDIS_URL cell found)")
        src_rd = None

    # --- EVM RPC (optional, from Alchemy) ----------------------------------
    if s.get("EVM_RPC_URL"):
        resolved["EVM_RPC_URL"] = s["EVM_RPC_URL"]
    elif evm_net and s.get("ALCHEMY_API_KEY"):
        resolved["EVM_RPC_URL"] = f"https://{evm_net}.g.alchemy.com/v2/{s['ALCHEMY_API_KEY']}"
    else:
        todos.append("EVM_RPC_URL (set EVM_NETWORK=eth-mainnet to build from ALCHEMY_API_KEY, "
                     "or add an EVM_RPC_URL row)")

    # --- Anvil fork source (read-only sim; PRIVATE_KEY stays empty in .mcp.json)
    if s.get("ANVIL_FORK_RPC_URL"):
        resolved["ANVIL_FORK_RPC_URL"] = s["ANVIL_FORK_RPC_URL"]
    elif resolved.get("EVM_RPC_URL"):
        resolved["ANVIL_FORK_RPC_URL"] = resolved["EVM_RPC_URL"]  # fork from the same Alchemy RPC
    else:
        todos.append("ANVIL_FORK_RPC_URL (add a row, or set EVM_NETWORK to derive from ALCHEMY_API_KEY)")

    # --- Pass-through optionals (only if the operator added a cell) ---------
    for var in ("GITHUB_TOKEN", "MAGIC_21ST_API_KEY", "CONTEXT7_API_KEY"):
        if s.get(var):
            resolved[var] = s[var]
        else:
            todos.append(f"{var} (add a '{var}' row to the Excel to wire it)")

    # --- Non-secret: sandbox root for filesystem MCP -----------------------
    resolved["PROJECT_ROOT"] = str(REPO_ROOT)

    # --- Write .env.mcp (gitignored) ---------------------------------------
    lines = [
        "# AUTO-GENERATED from ArbitrageX_Secrets_Config.xlsx by scripts/mcp/gen-env-mcp.py",
        "# DO NOT COMMIT. DO NOT HAND-EDIT. Edit the Excel SSOT and re-run the generator.",
        "# Doctrine: audit / shadow / read-only. No PRIVATE_KEY. Postgres = read-only role.",
        "",
    ]
    for k, v in resolved.items():
        lines.append(f"{k}={v}")
    if todos:
        lines.append("")
        lines.append("# --- NOT WIRED (add the matching row to the Excel, then re-run) ---")
        for t in todos:
            lines.append(f"# TODO {t}")
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
    try:
        os.chmod(OUT, 0o600)
    except Exception:
        pass

    # --- Masked report -----------------------------------------------------
    print(f"Source : {XLSX}")
    print(f"Output : {OUT}  (gitignored)")
    print(f"Discovered {len(s)} variable cell(s) across the workbook.")
    print("\nWired into .env.mcp:")
    for k, v in resolved.items():
        tag = "" if k == "PROJECT_ROOT" else f"  {mask(v)}"
        if k == "DATABASE_URL_READONLY":
            tag += f"   [{src_pg}]"
        if k == "REDIS_URL_READONLY":
            tag += f"   [{src_rd}]"
        print(f"  + {k}{tag}")
    if todos:
        print("\nNOT wired (operator action):")
        for t in todos:
            print(f"  - {t}")
    if warnings:
        print("\n!! DOCTRINE WARNINGS:")
        for w in warnings:
            print("  ! " + w)
    print("\nNext: load it and start Claude Code via the launcher:")
    print("  powershell -File scripts/mcp/run-claude-with-mcp.ps1")


if __name__ == "__main__":
    main()
