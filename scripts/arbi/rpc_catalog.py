#!/usr/bin/env python3
"""
ARBI — catálogo canónico de RPC + sonda de latencia.

Dos cosas, y las dos con artefacto:

  1) `extract`  — lee el Excel unificado (`ArbitrageX_Unified_Config.xlsm`, hoja
     "RPC Providers") y emite un catálogo JSON **versionable** con los RPC por
     cadena. Los secretos NO se copian: cualquier URL que lleve una API key
     embebida se sanea a `${VAR}` (ver `sanitize`). Un catálogo con credenciales
     dentro no es canónico, es una fuga esperando a pasar.

  2) `probe`    — mide LATENCIA REAL de cada RPC (`eth_chainId`, N muestras) y
     emite, por cadena, el orden de conexión: el más rápido primero, con su
     p50/p95. Eso es "engancharse siempre al mejor con menor latencia": no por
     fama del proveedor, por número medido.

R8 / RULE 00: un RPC que no responde se marca `ok=false` con su causa; nunca se le
asigna una latencia inventada.

RELACIÓN CON LO YA EXISTENTE (honestidad de instrumento, 2026-10-19):
  `scripts/arbx-env-deploy/gen_rpc_env_from_xlsx.py` YA emite
  `RPC_HTTP_<chain_id>` / `RPC_WS_<chain_id>` desde la MISMA hoja `RPC Providers`,
  en el formato exacto del `.env`, con un denylist de proveedores muertos.
  Es decir: el `extract` / `emit-env` / `apply` de este módulo es un SEGUNDO
  instrumento para lo mismo. El superviviente es `gen_rpc_env_from_xlsx.py`
  (ya está cableado en la macro y en el README). Lo que este módulo aporta y
  aquel NO tiene es la **sonda de latencia** (`probe` / `compare`): evidencia
  medida con ventana y control, frente a una lista declarada. La parte duplicada
  queda como candidata a retirada/plegado en la reconciliación de instrumentos.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import statistics
import sys
import time
import urllib.error
import urllib.request
from dataclasses import dataclass, asdict
from typing import Optional

UA = "arbi-rpc-catalog/1.0 (+https://github.com/hefarica/arbitragex-v2)"

# Cadenas con las que trabaja el pipeline (las mismas del selector).
WANTED_CHAINS = {
    "Ethereum Mainnet", "Base", "Arbitrum One", "Optimism", "Polygon Mainnet",
    "BSC Mainnet", "Avalanche", "Linea", "Scroll", "Blast", "Gnosis",
}

# Nombre de cadena en el Excel -> chain_id (la verdad de cadena vive en configs/app.toml).
CHAIN_IDS = {
    "Ethereum Mainnet": 1,
    "BSC Mainnet": 56,
    "Polygon Mainnet": 137,
    "Arbitrum One": 42161,
    "Optimism": 10,
    "Base": 8453,
    "Avalanche": 43114,
    "Gnosis": 100,
    "Linea": 59144,
    "Scroll": 534352,
    "Blast": 81457,
}

# Proveedor -> variable de entorno para la clave (el catálogo guarda `${VAR}`, nunca el valor).
PROVIDER_KEY_ENV = {
    "Alchemy": "ALCHEMY_API_KEY",
    "Infura": "INFURA_API_KEY",
    "QuickNode": "QUICKNODE_API_KEY",
}

_KEYLIKE = re.compile(r"(/v2/|/v3/|apikey=|/rpc/)[A-Za-z0-9_\-]{16,}", re.IGNORECASE)


def sanitize(url: str, provider: str) -> str:
    """Sustituye una API key embebida por `${VAR}`. No toca rutas públicas."""
    env = PROVIDER_KEY_ENV.get(provider)
    if "alchemy.com/v2/" in url:
        return re.sub(r"(/v2/)[^/?#]+", r"\1${ALCHEMY_API_KEY}", url)
    if env and _KEYLIKE.search(url):
        return _KEYLIKE.sub(lambda m: m.group(1) + "${" + env + "}", url)
    return url


def has_embedded_secret(url: str) -> bool:
    return "alchemy.com/v2/" in url or bool(_KEYLIKE.search(url))


# ─────────────────────────────── extracción ────────────────────────────────────

@dataclass
class RpcEntry:
    chain: str
    protocol: str
    provider: str
    url: str


def extract(xlsm_path: str, sheet: str = "RPC Providers") -> list[RpcEntry]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsm_path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"hoja '{sheet}' no está en el Excel; hay: {wb.sheetnames}")
    ws = wb[sheet]
    out: list[RpcEntry] = []
    for row in ws.iter_rows(values_only=True):
        c = [("" if v is None else str(v).strip()) for v in row]
        if len(c) >= 4 and c[0] and c[1] in ("HTTP", "WSS") and c[2] and c[3]:
            out.append(RpcEntry(chain=c[0], protocol=c[1], provider=c[2], url=c[3]))
    return out


def build_catalog(entries: list[RpcEntry]) -> dict:
    chains: dict[str, list[dict]] = {}
    secrets_found = 0
    for e in entries:
        if e.chain not in WANTED_CHAINS:
            continue
        if has_embedded_secret(e.url):
            secrets_found += 1
        chains.setdefault(e.chain, []).append({
            "protocol": e.protocol,
            "provider": e.provider,
            "url": sanitize(e.url, e.provider),
            "url_has_env_var": has_embedded_secret(e.url),
        })
    return {
        "_comment": ("Catálogo canónico de RPC por cadena. Extraído del Excel unificado; "
                     "las URLs con clave van saneadas a `${VAR}` y la clave vive en el `.env` del host, "
                     "NUNCA en este fichero."),
        "secrets_sanitized": secrets_found,
        "chains": {k: chains[k] for k in sorted(chains)},
    }


# ─────────────────────────── emisión G-RPC-1 (.env) ────────────────────────────

ENV_HEADER = [
    "# ── G-RPC-1 — catálogo de RPC (generado por scripts/arbi/rpc_catalog.py) ──",
    "# Formato: CSV de `nombre=url` por cadena. ≥2 proveedores por cadena (G-RPC-1);",
    "# con uno solo el pool no puede hacer failover y emite warn al arrancar.",
    "# Las claves van saneadas a ${VAR}: el valor vive en el almacén de secretos",
    "# (docs/operations/SECRETS_POLICY.md), NUNCA en este fichero.",
    "# Candidatos, no prioridades: el pool elige por latencia MEDIDA (EWMA).",
    "# Este fichero es IDEMPOTENTE: aplicarlo dos veces da el mismo byte.",
    "",
]

_ENV_LINE = re.compile(r"^RPC_(HTTP|WSS)_(\d+)=(.*)$")


def model_from_catalog(catalog: dict) -> dict[int, dict[str, dict[str, str]]]:
    """Excel/catálogo -> modelo interno `{chain_id: {'HTTP': {nombre: url}, 'WSS': {...}}}`."""
    model: dict[int, dict[str, dict[str, str]]] = {}
    for chain, entries in catalog.get("chains", {}).items():
        cid = CHAIN_IDS.get(chain)
        if cid is None:
            continue
        slot = model.setdefault(cid, {"HTTP": {}, "WSS": {}})
        for e in entries:
            name = e["provider"].lower().replace(" ", "")
            slot.setdefault(e["protocol"], {})[name] = e["url"]
    return model


def model_from_env(text: str) -> dict[int, dict[str, dict[str, str]]]:
    """G-RPC-1 `.env` -> modelo interno. Entrada POLIMÓRFICA: acepta `nombre=url` y
    una URL desnuda (se nombra `primary`, como el pool)."""
    model: dict[int, dict[str, dict[str, str]]] = {}
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        m = _ENV_LINE.match(line)
        if not m:
            continue
        proto, cid, rest = m.group(1), int(m.group(2)), m.group(3).strip()
        if not rest:
            continue
        slot = model.setdefault(cid, {"HTTP": {}, "WSS": {}})
        for tok in rest.split(","):
            tok = tok.strip()
            if not tok:
                continue
            name, url = (tok.split("=", 1) if "=" in tok else ("primary", tok))
            slot.setdefault(proto, {})[name.strip().lower()] = url.strip()
    return model


def emit_env_from_model(model: dict[int, dict[str, dict[str, str]]]) -> str:
    """Modelo -> texto G-RPC-1, DETERMINISTA (claves ordenadas) => idempotente."""
    lines = list(ENV_HEADER)
    for cid in sorted(model):
        for proto in ("HTTP", "WSS"):
            d = model[cid].get(proto) or {}
            if not d:
                continue
            toks = ",".join(f"{k}={d[k]}" for k in sorted(d))
            lines.append(f"RPC_{proto}_{cid}={toks}")
        lines.append("")
    return "\n".join(lines) + "\n"


def validate_model(model: dict[int, dict[str, dict[str, str]]], min_providers: int = 2) -> list[str]:
    """Fail-closed: devuelve la lista de errores. Vacía = se puede aplicar."""
    errs: list[str] = []
    for cid in sorted(model):
        for proto in ("HTTP", "WSS"):
            d = model[cid].get(proto) or {}
            if not d:
                continue
            if len(d) < min_providers:
                errs.append(f"chain {cid} {proto}: {len(d)} proveedor(es) < {min_providers} — failover imposible")
            for name, url in sorted(d.items()):
                if not re.match(r"^(https?|wss?)://", url):
                    errs.append(f"chain {cid} {proto} {name}: URL no válida: {url}")
                elif "${" not in url and has_embedded_secret(url):
                    errs.append(f"chain {cid} {proto} {name}: URL con clave embebida (fuga): {url}")
    return errs


def emit_env(catalog: dict, min_providers: int = 2) -> tuple[str, list[str]]:
    """Compatibilidad: catálogo -> (texto G-RPC-1, avisos de validación)."""
    model = model_from_catalog(catalog)
    return emit_env_from_model(model), validate_model(model, min_providers)


# ─────────────────────────────── sonda de latencia ─────────────────────────────

def _rpc_timed(url: str, method: str = "eth_chainId", timeout: float = 8.0) -> tuple[bool, float, str]:
    body = json.dumps({"jsonrpc": "2.0", "id": 1, "method": method, "params": []}).encode()
    req = urllib.request.Request(url, data=body,
                                 headers={"content-type": "application/json",
                                          "accept": "application/json", "user-agent": UA})
    t0 = time.perf_counter()
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            d = json.loads(r.read().decode("utf-8"))
        dt = (time.perf_counter() - t0) * 1000.0
        if d.get("result") is not None:
            return True, dt, ""
        return False, dt, str(d.get("error"))[:120]
    except Exception as exc:  # noqa: BLE001
        return False, (time.perf_counter() - t0) * 1000.0, f"{type(exc).__name__}: {exc}"[:120]


def percentile(sorted_vals: list[float], p: float) -> float:
    if not sorted_vals:
        return float("nan")
    idx = min(len(sorted_vals) - 1, int(round(p * (len(sorted_vals) - 1))))
    return sorted_vals[idx]


def probe_url(url: str, samples: int = 5, timeout: float = 8.0) -> dict:
    """Mide un endpoint. Declara SIEMPRE n (muestras ok / total): un p50 sin n no existe."""
    lat: list[float] = []
    last_err = ""
    for _ in range(samples):
        ok, dt, err = _rpc_timed(url, timeout=timeout)
        if ok:
            lat.append(dt)
        else:
            last_err = err
    if not lat:
        return {"ok": False, "n_ok": 0, "n_total": samples,
                "p50_ms": None, "p95_ms": None, "p99_ms": None, "min_ms": None, "max_ms": None,
                "error": last_err or "sin respuesta"}
    s = sorted(lat)
    return {
        "ok": True, "n_ok": len(s), "n_total": samples,
        "p50_ms": round(statistics.median(s), 1),
        "p95_ms": round(percentile(s, 0.95), 1),
        "p99_ms": round(percentile(s, 0.99), 1),
        "min_ms": round(s[0], 1), "max_ms": round(s[-1], 1),
        "error": "",
    }


def run_controls(good_url: str, dead_url: str, samples: int, timeout: float) -> dict:
    """CONTROL de la sonda: un endpoint que sabemos bueno y uno que sabemos muerto.
    Si la sonda no los distingue, el instrumento está roto y ningún p50 vale."""
    g = probe_url(good_url, samples, timeout)
    d = probe_url(dead_url, samples, timeout)
    ok = bool(g["ok"]) and not bool(d["ok"])
    return {
        "known_good": {"url": good_url, **g},
        "known_dead": {"url": dead_url, **d},
        "distinguishes": ok,
        "verdict": ("la sonda distingue bueno de muerto" if ok
                    else "CONTROL FALLIDO: la sonda no distingue bueno de muerto — sus números no valen"),
    }


def probe_catalog(catalog: dict, chains: Optional[list[str]], samples: int, timeout: float,
                  control_good: str, control_dead: str) -> dict:
    started = time.time()
    controls = run_controls(control_good, control_dead, samples, timeout)
    result: dict[str, list[dict]] = {}
    for chain, entries in catalog.get("chains", {}).items():
        if chains and chain not in chains:
            continue
        scored = []
        for e in entries:
            if e["protocol"] != "HTTP":
                continue  # la sonda mide HTTP; WSS se ordena aparte, sin fingir
            if e.get("url_has_env_var"):
                scored.append({**e, "ok": None, "n_ok": 0, "n_total": 0,
                               "p50_ms": None, "p95_ms": None, "p99_ms": None,
                               "error": "requiere ${VAR}: no se sondea sin clave"})
                continue
            m = probe_url(e["url"], samples, timeout)
            scored.append({**e, **m})
        # orden: los medidos por p50 ascendente; al final los no medidos
        scored.sort(key=lambda s: (s["p50_ms"] is None, s["p50_ms"] if s["p50_ms"] is not None else 1e9))
        for i, s in enumerate(scored):
            s["rank"] = i + 1
            # EVIDENCIA, no selector: el HttpRpcPool elige por EWMA. Esto sólo dice
            # quién fue el más rápido EN ESTA VENTANA.
            s["use"] = ("fastest_measured" if i == 0 and s["p50_ms"] is not None
                        else ("measured" if s["p50_ms"] is not None else "unprobed"))
        result[chain] = scored
    ended = time.time()
    return {
        "_comment": ("EVIDENCIA de latencia, NO selector. El HttpRpcPool "
                     "(backend/shared-rs/src/rpc_failover.rs) elige por EWMA + circuit breaker; "
                     "este orden es el de ESTA ventana y no debe cablearse como prioridad."),
        "samples_per_endpoint": samples,
        "window": {
            "started_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(started)),
            "ended_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(ended)),
            "duration_s": round(ended - started, 1),
        },
        "controls": controls,
        "chains": result,
    }


def compare_runs(a: dict, b: dict) -> dict:
    """Criterio de muerte: si dos corridas dan ordenamientos distintos sin cambio de red,
    la sonda mide RUIDO y va antes que cualquier ajuste."""
    out: dict = {"chains": {}, "unstable": []}
    for chain in sorted(set(a.get("chains", {})) & set(b.get("chains", {}))):
        order_a = [e["provider"] for e in a["chains"][chain] if e.get("p50_ms") is not None]
        order_b = [e["provider"] for e in b["chains"][chain] if e.get("p50_ms") is not None]
        same = order_a == order_b
        out["chains"][chain] = {"a": order_a, "b": order_b, "same_order": same}
        if not same:
            out["unstable"].append(chain)
    out["verdict"] = ("estable: mismo ordenamiento en las dos corridas" if not out["unstable"]
                      else f"INESTABLE en {out['unstable']} — posible ruido; va antes que cualquier ajuste")
    return out


def main() -> int:
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="ARBI RPC catalog + latency probe")
    sub = ap.add_subparsers(dest="cmd", required=True)

    pe = sub.add_parser("extract", help="Excel -> catálogo canónico saneado")
    pe.add_argument("--xlsm", required=True)
    pe.add_argument("--out", default=os.path.join(here, "rpc_catalog.json"))

    pv = sub.add_parser("emit-env", help="catálogo -> bloque .env en formato G-RPC-1")
    pv.add_argument("--catalog", default=os.path.join(here, "rpc_catalog.json"))
    pv.add_argument("--out", default="rpc_catalog.env")
    pv.add_argument("--min-providers", type=int, default=2)

    pp = sub.add_parser("probe", help="catálogo -> latencias medidas y orden de conexión")
    pp.add_argument("--catalog", default=os.path.join(here, "rpc_catalog.json"))
    pp.add_argument("--chain", action="append", default=None)
    pp.add_argument("--samples", type=int, default=5)
    pp.add_argument("--timeout", type=float, default=8.0)
    pp.add_argument("--control-good", default="https://ethereum-rpc.publicnode.com/")
    pp.add_argument("--control-dead", default="https://eth.llamarpc.com")
    pp.add_argument("--out", default="rpc_latency.json")

    pa = sub.add_parser("apply", help="normaliza desde xlsm/env/json, valida y emite G-RPC-1 (idempotente)")
    pa.add_argument("--from", dest="src", required=True)
    pa.add_argument("--format", choices=["auto", "xlsm", "env", "json"], default="auto")
    pa.add_argument("--out", default="rpc_catalog.env")
    pa.add_argument("--min-providers", type=int, default=2)

    pval = sub.add_parser("validate", help="valida un .env G-RPC-1 (fail-closed)")
    pval.add_argument("--file", required=True)
    pval.add_argument("--min-providers", type=int, default=2)

    pc = sub.add_parser("compare", help="dos corridas -> estabilidad del ordenamiento (criterio de muerte)")
    pc.add_argument("--a", required=True)
    pc.add_argument("--b", required=True)
    pc.add_argument("--out", default="rpc_stability.json")

    args = ap.parse_args()

    if args.cmd == "extract":
        entries = extract(args.xlsm)
        cat = build_catalog(entries)
        with open(args.out, "w", encoding="utf-8") as fh:
            json.dump(cat, fh, indent=2, ensure_ascii=False)
        total = sum(len(v) for v in cat["chains"].values())
        print(f"cadenas={len(cat['chains'])} rpcs={total} secretos_saneados={cat['secrets_sanitized']} -> {args.out}")
        return 0

    if args.cmd == "emit-env":
        cat = json.load(open(args.catalog, encoding="utf-8"))
        text, warnings = emit_env(cat, args.min_providers)
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(text)
        for w in warnings:
            print(f"::warning:: {w}")
        print(text)
        print(f"-> {args.out}  (avisos: {len(warnings)})")
        return 0

    if args.cmd == "apply":
        fmt = args.format
        if fmt == "auto":
            ext = os.path.splitext(args.src)[1].lower()
            fmt = {"xlsm": "xlsm", "json": "json"}.get(ext.lstrip("."), "env")
        if fmt == "xlsm":
            model = model_from_catalog(build_catalog(extract(args.src)))
        elif fmt == "json":
            model = model_from_catalog(json.load(open(args.src, encoding="utf-8")))
        else:
            model = model_from_env(open(args.src, encoding="utf-8").read())
        errs = validate_model(model, args.min_providers)
        if errs:
            for e in errs:
                print(f"::error:: {e}")
            print(f"NO SE APLICA: {len(errs)} error(es) de validación (fail-closed)")
            return 1
        text = emit_env_from_model(model)
        if os.path.exists(args.out) and open(args.out, encoding="utf-8").read() == text:
            print(f"sin cambios — {args.out} ya es idéntico (idempotente)")
        else:
            with open(args.out, "w", encoding="utf-8") as fh:
                fh.write(text)
            print(f"escrito {args.out} ({len(model)} cadenas)")
        return 0

    if args.cmd == "validate":
        model = model_from_env(open(args.file, encoding="utf-8").read())
        errs = validate_model(model, args.min_providers)
        for e in errs:
            print(f"::error:: {e}")
        if errs:
            print(f"INVÁLIDO: {len(errs)} error(es)")
            return 1
        print(f"VÁLIDO: {len(model)} cadenas, ≥{args.min_providers} proveedores por cadena, sin claves embebidas")
        return 0

    if args.cmd == "compare":
        a = json.load(open(args.a, encoding="utf-8"))
        b = json.load(open(args.b, encoding="utf-8"))
        cmp = compare_runs(a, b)
        with open(args.out, "w", encoding="utf-8") as fh:
            json.dump(cmp, fh, indent=2, ensure_ascii=False)
        for chain, d in cmp["chains"].items():
            mark = "OK " if d["same_order"] else "DIFF"
            print(f"{mark} {chain}: {d['a']} vs {d['b']}")
        print(cmp["verdict"])
        return 0

    cat = json.load(open(args.catalog, encoding="utf-8"))
    res = probe_catalog(cat, args.chain, args.samples, args.timeout, args.control_good, args.control_dead)
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(res, fh, indent=2, ensure_ascii=False)
    ctl = res["controls"]
    print(f"CONTROLES: bueno={ctl['known_good']['ok']} muerto={ctl['known_dead']['ok']} -> {ctl['verdict']}")
    print(f"VENTANA: {res['window']['started_at']} .. {res['window']['ended_at']} "
          f"({res['window']['duration_s']}s) · n={res['samples_per_endpoint']} por endpoint")
    for chain, entries in res["chains"].items():
        print(f"\n== {chain}")
        for e in entries:
            if e["p50_ms"] is None:
                print(f"   - {e['provider']:<11} {e['use']:<16} n={e['n_ok']}/{e['n_total']} ({e['error']})")
            else:
                print(f"   {e['rank']:>2}. {e['provider']:<11} {e['use']:<16} "
                      f"n={e['n_ok']}/{e['n_total']} p50={e['p50_ms']:>7.1f} p95={e['p95_ms']:>7.1f} p99={e['p99_ms']:>7.1f} ms")
    return 0


if __name__ == "__main__":
    sys.exit(main())
